"""Extrait en lecture seule les tarifs des trois tableurs fournis."""
import json
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("P:/formation/drop/GAEC/TARIFS GAEC")
OUT = Path(__file__).resolve().parent.parent / "data" / "tariff-reference.json"
NS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
}


def price(value):
    try:
        number = float(str(value).replace(",", "."))
        return round(number, 4) if number > 0 else None
    except (TypeError, ValueError):
        return None


def ods_sheets(file):
    with zipfile.ZipFile(file) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    for table in root.findall(".//table:table", NS):
        rows = []
        for row in table.findall("table:table-row", NS):
            cells = []
            for cell in row:
                if cell.tag not in {
                    f"{{{NS['table']}}}table-cell",
                    f"{{{NS['table']}}}covered-table-cell",
                }:
                    continue
                value = cell.get(f"{{{NS['office']}}}value")
                cells.append(value if value is not None else "".join(cell.itertext()).strip())
            rows.append(cells)
        yield table.get(f"{{{NS['table']}}}name"), rows


def add(rows, name, unit, amount, source):
    if not str(name or "").strip() or price(amount) is None:
        return
    rows.append({
        "name": str(name).strip(),
        "unit": str(unit or "").strip(),
        "price": price(amount),
        "source": source,
    })


def main():
    result = {"Mercuriale 2026": [], "Tarif épicerie": [], "Paniers": []}

    file = BASE / "MERCURIALE 2026 LA RAVOIRE (1).xlsx"
    sheet = load_workbook(file, read_only=True, data_only=True).active
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_number <= 2:
            continue
        add(result["Mercuriale 2026"], row[0], "", row[1], f"{file.name}:Feuil1:{row_number}")

    file = BASE / "tarifs épicerie.ods"
    for sheet_name, rows in ods_sheets(file):
        for row_number, row in enumerate(rows, 1):
            if row_number <= 2 or len(row) < 3:
                continue
            add(result["Tarif épicerie"], row[0], row[2], row[1], f"{file.name}:{sheet_name}:{row_number}")

    file = BASE / "tarifs vente directe.ods"
    for sheet_name, rows in ods_sheets(file):
        for row_number, row in enumerate(rows, 1):
            if sheet_name == "Feuille1":
                if row_number == 1 or len(row) < 3:
                    continue
                name, amount, unit = row[0:3]
            else:
                if len(row) < 4:
                    continue
                name, amount, unit = row[1:4]
            add(result["Paniers"], name, unit, amount, f"{file.name}:{sheet_name}:{row_number}")

    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name, rows in result.items():
        print(f"{name}: {len(rows)} lignes tarifées")
    print(f"Référence créée : {OUT}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Erreur : {error}", file=sys.stderr)
        sys.exit(1)
